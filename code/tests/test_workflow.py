"""Regression tests for notebook orchestration; run separately from starter tests."""
import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path

import numpy as np
import pandas as pd
import torch
from PIL import Image

CODE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(CODE))
sys.path.insert(1, str(CODE.parent))
# Avoid the Windows PyArrow native crash while testing the remaining contracts.
pd.options.mode.string_storage = "python"

import dataset
import inference
import train
import model as models


class WorkflowTests(unittest.TestCase):
    def workflow(self):
        spec = importlib.util.spec_from_file_location("lab_workflow", CODE / "lab_workflow.py")
        if not spec.origin or not Path(spec.origin).exists():
            self.fail("Notebook workflow has not been implemented")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module

    def test_fivecrop_center_and_corners(self):
        x = torch.arange(16).reshape(1, 1, 4, 4)
        crops = inference.views_multicrop(x, 2)
        self.assertEqual([c.flatten().tolist() for c in crops],
                         [[0, 1, 4, 5], [2, 3, 6, 7], [8, 9, 12, 13],
                          [10, 11, 14, 15], [5, 6, 9, 10]])

    def test_gmac_counter_uses_model_device_and_dtype(self):
        model = torch.nn.Conv2d(3, 2, kernel_size=3).double().train()
        # 6x6 outputs, two channels, 27 multiply-accumulates each.
        self.assertAlmostEqual(models.count_gmacs(model, 8), 1944 / 1e9)
        self.assertTrue(model.training)
        if torch.cuda.is_available():
            self.assertAlmostEqual(models.count_gmacs(model.cuda(), 8), 1944 / 1e9)

    def test_gmac_counter_counts_every_token_in_linear_layers(self):
        model = torch.nn.Sequential(torch.nn.Flatten(2), torch.nn.Linear(64, 2))
        self.assertAlmostEqual(models.count_gmacs(model, 8), 384 / 1e9, places=12)

    def test_cache_preserves_original_images_when_training_size_differs(self):
        with tempfile.TemporaryDirectory() as tmp:
            Image.fromarray(np.full((64, 64, 3), 17, np.uint8)).save(Path(tmp) / "a.png")
            df = pd.DataFrame({"Filename": ["a.png"], "Label": [0]})
            ds = dataset.DeepWeedsDataset(df, tmp, cache_in_ram=True, img_size=32)
            self.assertEqual(ds[0][0].size, (64, 64))
            self.assertEqual(np.asarray(ds[0][0])[0, 0].tolist(), [17, 17, 17])

    def test_debug_run_cannot_evaluate_test(self):
        self.assertIn("debug_samples_per_class", train.Config.__dataclass_fields__,
                      "No isolated smoke-run configuration exists")
        cfg = train.Config(debug_samples_per_class=2, save_test_predictions=True)
        with self.assertRaisesRegex(ValueError, "debug.*test|test.*debug"):
            train.run(cfg)

    def test_local_profile_finds_images_and_conserves_memory(self):
        flow = self.workflow()
        cfg = flow.base_config(CODE.parent, "local")
        self.assertEqual(Path(cfg.images_dir), CODE.parent / "images")
        self.assertLessEqual(cfg.batch_size, 8)
        self.assertEqual(cfg.num_workers, 0)
        self.assertFalse(cfg.save_test_predictions)

    def test_selection_rejects_scores_from_a_different_recipe(self):
        flow = self.workflow()
        import dataclasses
        with tempfile.TemporaryDirectory() as tmp:
            cfg = flow.base_config(tmp)
            cfg.exp_id = "B01"
            path = train.run_dir(cfg)
            path.mkdir(parents=True)
            (path / "summary.json").write_text(json.dumps({"val_macro_f1": .8}))
            old = dataclasses.asdict(cfg)
            old["epochs"] = 1
            (path / "config.json").write_text(json.dumps(old))
            with self.assertRaisesRegex(ValueError, "Config|config"):
                flow.best_config([cfg])

    def test_amp_overflow_does_not_advance_lr_schedule(self):
        model = torch.nn.Linear(3, 9)
        optimizer = torch.optim.SGD(model.parameters(), lr=.1)
        scheduler = torch.optim.lr_scheduler.StepLR(optimizer, step_size=1)
        scaler = torch.amp.GradScaler("cpu")
        cfg = train.Config(amp=False)
        def nonfinite(logits, y):
            return logits.sum() * float("nan")
        train.train_one_epoch(model, [(torch.ones(2, 3), torch.zeros(2, dtype=torch.long), ['a', 'b'])],
                              nonfinite, optimizer, scheduler, scaler, cfg, torch.device('cpu'))
        self.assertEqual(scheduler.last_epoch, 0)
        self.assertEqual(optimizer.param_groups[0]['lr'], .1)

    def test_final_recovery_rebuilds_metrics_without_model_evaluation(self):
        flow = self.workflow()
        import eval as ev
        with tempfile.TemporaryDirectory() as tmp:
            cfg = flow.base_config(tmp)
            cfg.exp_id = 'F01'
            d = train.run_dir(cfg)
            d.mkdir(parents=True)
            (d / 'summary.json').write_text(json.dumps({'exp_id': 'F01', 'seed': 0}))
            (d / 'final_inference.json').write_text(json.dumps({'method': 'single', 'temperature': 1.}))
            probs = np.eye(9)
            for split in ('val', 'test'):
                ev.save_predictions(Path(cfg.pred_dir) / f'F01_seed0_{split}.csv',
                                    [f'{i}.jpg' for i in range(9)], np.arange(9), probs)
            self.assertTrue(hasattr(flow, 'recover_final'), 'Cannot repair an interrupted final write')
            flow.recover_final(cfg)
            summary = json.loads((d / 'summary.json').read_text())
            self.assertEqual(summary['test_macro_f1'], 1.)
            self.assertEqual(summary['val_macro_f1'], 1.)

    def test_hflip_probability_average_preserves_names(self):
        flow = self.workflow()
        class FirstPixel(torch.nn.Module):
            def forward(self, x):
                logits = torch.zeros(x.shape[0], 9)
                logits[:, 0] = x[:, 0, 0, 0]
                return logits
        images = torch.tensor([[[[0., 2.]]], [[[2., 0.]]]])
        names, y, probs = flow.predict_method(FirstPixel(),
            [(images, torch.tensor([1, 0]), ['first.jpg', 'second.jpg'])], torch.device('cpu'), 'hflip_prob')
        self.assertEqual(names, ['first.jpg', 'second.jpg'])
        np.testing.assert_array_equal(y, [1, 0])
        p0 = .5 * (1/9 + np.exp(2)/(np.exp(2)+8))
        p_other = .5 * (1/9 + 1/(np.exp(2)+8))
        np.testing.assert_allclose(probs, [[p0] + [p_other]*8]*2, atol=1e-7)

    def test_baseline_recovery_after_csv_before_summary(self):
        flow = self.workflow()
        import eval as ev
        with tempfile.TemporaryDirectory() as tmp:
            cfg = flow.base_config(tmp)
            d = train.run_dir(cfg)
            d.mkdir(parents=True)
            for split in ('val', 'test'):
                ev.save_predictions(Path(cfg.pred_dir) / f'T00_seed0_{split}.csv',
                                    [f'{i}.jpg' for i in range(9)], np.arange(9), np.eye(9))
            self.assertTrue(hasattr(flow, 'recover_baseline'), 'Cannot repair baseline test metrics')
            flow.recover_baseline(cfg)
            summary = json.loads((d / 'summary.json').read_text())
            self.assertEqual(summary['test_macro_f1'], 1.)

    def test_export_excludes_debug_and_uses_sample_std(self):
        flow = self.workflow()
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for exp, seed, f1, debug in [("F01", 0, .7, False), ("F01", 1, .9, False),
                                          ("SMOKE", 0, 1., True)]:
                d = root / "runs" / exp / f"seed{seed}"
                d.mkdir(parents=True)
                (d / "summary.json").write_text(json.dumps({
                    "exp_id": exp, "seed": seed, "backbone": "tiny", "debug_run": debug,
                    "val_macro_f1": f1, "test_macro_f1": f1, "test_top1": f1,
                    "test_ece": .1}), encoding="utf-8")
            out = flow.export_results(root)
            book = load_workbook(out, data_only=True)
            self.assertEqual(set(book.sheetnames), {"Summary", "Backbones", "Training", "Inference",
                                                   "Final", "PerClass", "Latency"})
            summary_rows = list(book["Summary"].values)
            self.assertNotIn("SMOKE", str(summary_rows))
            rows = list(book["Final"].values)
            header = rows[0]
            agg = next(dict(zip(header, row)) for row in rows[1:] if row[header.index("seed")] == "mean_std")
            self.assertAlmostEqual(agg["test_macro_f1_mean"], .8)
            self.assertAlmostEqual(agg["test_macro_f1_std"], .1414213562373095)
            book.close()


if __name__ == "__main__":
    unittest.main()
