"""Runnable notebook workflow built on the original lab modules.

All selection uses validation. Smoke runs and official runs use separate folders.
The exporter reads saved results; missing experiments remain missing.
"""
from __future__ import annotations

import runtime
import argparse
import dataclasses
import json
from pathlib import Path

import numpy as np
import pandas as pd
import torch
import torch.nn.functional as F

import dataset as ds
import inference as inf
from benchmark import bench
from train import Config, _import_eval, load_state, run, run_dir, set_seed

METHODS = ("single", "hflip_prob", "fivecrop_prob", "hflip_logit", "temperature")


def base_config(root, profile="local"):
    root = Path(root).resolve()
    candidates = (root / "data/images", root / "images")
    images = next((p for p in candidates if p.is_dir()), candidates[0])
    if profile not in ("local", "kaggle"):
        raise ValueError("profile must be local or kaggle")
    return Config(images_dir=str(images), labels_dir=str(root / "data/labels"),
                  out_dir=str(root / "runs"), pred_dir=str(root / "predictions"),
                  curves_dir=str(root / "curves"), logit_dir=str(root / "logits"),
                  batch_size=4 if profile == "local" else 32,
                  num_workers=0 if profile == "local" else 2,
                  cache_images=False, pin_memory=False if profile == "local" else None)


def experiment_plan(base, stage):
    if stage == "backbones":
        names = ("resnet50", "convnext_tiny", "deit_small", "swin_tiny", "mobilenetv3")
        return [dataclasses.replace(base, exp_id=f"B{i:02}", backbone=name,
                                    save_test_predictions=False) for i, name in enumerate(names, 1)]
    if stage == "training":
        variants = [({}, "T00"), ({"init": "scratch"}, "T01"),
                    ({"init": "frozen"}, "T02"), ({"aug": "trivial"}, "T03"),
                    ({"aug": "color"}, "T04"),
                    ({"loss": "ls", "label_smoothing": .1}, "T05"),
                    ({"loss": "focal"}, "T06"), ({"mix": "cutmix"}, "T07"),
                    ({"mix": "mixup"}, "T08")]
        return [dataclasses.replace(base, exp_id=tag, save_test_predictions=False, **kw)
                for kw, tag in variants]
    raise ValueError("stage must be backbones or training")


def run_experiments(configs):
    rows = []
    for cfg in configs:
        if cfg.save_test_predictions or cfg.debug_samples_per_class is not None:
            raise ValueError("selection experiments must use full train/val, without test")
        summary = run_dir(cfg) / "summary.json"
        if summary.exists():
            validate_saved_config(cfg)
            rows.append(json.loads(summary.read_text(encoding="utf-8")))
            print(f"Reuse completed run: {cfg.exp_id}, seed {cfg.seed}")
        else:
            rows.append(run(cfg))
        if torch.cuda.is_available():
            torch.cuda.empty_cache()
    return pd.DataFrame(rows)


def validate_saved_config(cfg):
    old = json.loads((run_dir(cfg) / "config.json").read_text(encoding="utf-8"))
    current = dataclasses.asdict(cfg)
    # These only control writing/test access and resuming, not the scientific recipe.
    for key in ("save_test_predictions", "resume"):
        old.pop(key, None)
        current.pop(key, None)
    if old != current:
        raise ValueError(f"Config changed for {cfg.exp_id}; choose a new exp_id")


def best_config(configs):
    scores = []
    for cfg in configs:
        p = run_dir(cfg) / "summary.json"
        if not p.exists():
            raise FileNotFoundError(f"Complete all comparisons first: {p}")
        validate_saved_config(cfg)
        result = json.loads(p.read_text(encoding="utf-8"))
        if result.get("debug_run"):
            raise ValueError("Cannot select a debug run")
        scores.append(result["val_macro_f1"])
    return configs[int(np.argmax(scores))]


def _view_logits(model, x, method):
    if method in ("single", "temperature"):
        return [model(x)]
    if method.startswith("hflip"):
        return [model(x), model(inf.view_hflip(x))]
    if method == "fivecrop_prob":
        size = x.shape[-1]
        crops = inf.views_multicrop(x, max(1, size * 7 // 8))
        return [model(F.interpolate(c, size=(size, size), mode="bilinear", align_corners=False))
                for c in crops]
    raise ValueError(f"Unknown inference method: {method}")


def _aggregate(logits, method, temperature):
    if method == "hflip_logit":
        return torch.stack(logits).mean(0).softmax(-1)
    return torch.stack([(z / temperature if method == "temperature" else z).softmax(-1)
                        for z in logits]).mean(0)


def predict_method(model, loader, device, method="single", temperature=1., return_logits=False):
    if method not in METHODS:
        raise ValueError(f"Unknown inference method: {method}")
    model.eval()
    if return_logits and method not in ("single", "temperature"):
        raise ValueError("Raw logits are only available for single-view methods")
    names, targets, outputs, raw = [], [], [], []
    with torch.inference_mode():
        for images, labels, files in loader:
            logits = _view_logits(model, images.to(device), method)
            probs = _aggregate(logits, method, temperature).float().cpu().numpy()
            if return_logits:
                raw.append(logits[0].float().cpu().numpy())
            names.extend(files)
            targets.append(labels.numpy())
            outputs.append(probs)
    if not outputs:
        raise ValueError("Cannot evaluate an empty split")
    result = names, np.concatenate(targets), np.concatenate(outputs)
    return (*result, np.concatenate(raw)) if return_logits else result


def _loader(cfg, frame):
    return ds.make_loader(frame, cfg.images_dir, ds.build_transforms(False, cfg.img_size),
                          cfg.batch_size, train=False, num_workers=cfg.num_workers,
                          pin_memory=cfg.pin_memory, cache_in_ram=False)


def measure_method(model, cfg, device, method, temperature=1., iters=50):
    """Time actual views + aggregation on device, excluding disk I/O and decoding."""
    model.eval()
    x = torch.randn(1, 3, cfg.img_size, cfg.img_size, device=device)
    def forward():
        with torch.inference_mode():
            logits = _view_logits(model, x, method)
            return _aggregate(logits, method, temperature)
    timing = bench(forward, warmup=10, iters=iters,
                   sync=torch.cuda.synchronize if device.type == "cuda" else None)
    timing.update(gpu=torch.cuda.get_device_name(0) if device.type == "cuda" else "CPU",
                  dtype="fp32", batch=1, img_size=cfg.img_size,
                  images_per_s=1000 / timing["p50"], includes_preprocess=True,
                  timing_scope="tensor views + forward + aggregation; excludes image decoding")
    return timing


def compare_inference(cfg):
    if cfg.debug_samples_per_class is not None:
        raise ValueError("Inference selection cannot use debug data")
    validate_saved_config(cfg)
    ev = _import_eval()
    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    model, _, _ = load_state(cfg, device)
    _, val, _ = ds.load_split(cfg.labels_dir, cfg.fold)
    loader = _loader(cfg, val)
    _, y, logits = inf.predict_logits(model, loader, device)
    temperature = inf.fit_temperature(logits, y)
    rows = []
    out = Path(cfg.out_dir) / "inference" / cfg.exp_id / f"seed{cfg.seed}"
    out.mkdir(parents=True, exist_ok=True)
    for index, method in enumerate(METHODS):
        names, y, probs = predict_method(model, loader, device, method, temperature)
        m = ev.compute_metrics(y, probs.argmax(1), probs)
        row = {"exp_id": f"I{index:02}", "source_exp_id": cfg.exp_id,
               "seed": cfg.seed, "backbone": cfg.backbone, "method": method,
               "k_views": 5 if method == "fivecrop_prob" else 2 if method.startswith("hflip") else 1,
               "temperature": temperature if method == "temperature" else 1.,
               "val_macro_f1": m["macro_f1"], "val_top1": m["top1"], "val_ece": m["ece"],
               "checkpoint": str(run_dir(cfg) / "best.pt"),
               "latency": measure_method(model, cfg, device, method, temperature)}
        row["relative_cost"] = row["latency"]["p50"] / (rows[0]["latency"]["p50"] if rows else row["latency"]["p50"])
        ev.save_predictions(out / f"{row['exp_id']}_seed{cfg.seed}_val.csv", names, y, probs)
        (out / f"{row['exp_id']}.json").write_text(json.dumps(row, indent=2), encoding="utf-8")
        rows.append(row)
        print(method, f"val macro-F1={m['macro_f1']:.4f}", f"p95={row['latency']['p95']:.2f} ms")
    return pd.DataFrame([{k: v for k, v in r.items() if k != "latency"} for r in rows])


def lock_selection(cfg, method):
    if method not in METHODS or cfg.debug_samples_per_class is not None:
        raise ValueError("Select an official configuration and supported inference method")
    validate_saved_config(cfg)
    cfg = dataclasses.replace(cfg, save_test_predictions=False, resume=False)
    selected = {"config": dataclasses.asdict(cfg), "method": method,
                "selection_split": "val"}
    path = Path(cfg.out_dir) / "selection.json"
    if path.exists() and json.loads(path.read_text(encoding="utf-8")) != selected:
        raise ValueError("Final configuration already locked. Use a separate output directory for a new study.")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(selected, indent=2), encoding="utf-8")
    return path


def recover_final(cfg):
    """Repair interrupted final artifacts using saved predictions, never test forward."""
    ev = _import_eval()
    d = run_dir(cfg)
    metadata = json.loads((d / "final_inference.json").read_text(encoding="utf-8"))
    fp = Path(cfg.pred_dir) / f"F01_seed{cfg.seed}_test.csv"
    vp = Path(cfg.pred_dir) / f"F01_seed{cfg.seed}_val.csv"
    pending = d / "final_test_arrays.npz"
    if not fp.exists():
        # Saved immediately after the only test forward; recover a missing CSV.
        with np.load(pending, allow_pickle=True) as data:
            ev.save_predictions(fp, data['names'].tolist(), data['y'], data['probs'])
    if metadata['method'] == 'temperature':
        uncal = Path(cfg.pred_dir) / f"F01uncal_seed{cfg.seed}_test.csv"
        if not uncal.exists():
            with np.load(pending, allow_pickle=True) as data:
                ev.save_predictions(uncal, data['names'].tolist(), data['y'],
                                    inf.apply_temperature(data['logits'], 1.))
    summary_path = d / "summary.json"
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    summary.update(metadata)
    for split, path in (("val", vp), ("test", fp)):
        pred = ev.read_pred(str(path))
        metric = ev.compute_metrics(pred.y_true, pred.y_pred, pred.probs)
        summary.update({f"{split}_{key}": metric[key] for key in ("macro_f1", "top1", "ece", "balanced_acc")})
    summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")


def recover_baseline(cfg):
    ev = _import_eval()
    summary_path = run_dir(cfg) / "summary.json"
    summary = (json.loads(summary_path.read_text(encoding="utf-8")) if summary_path.exists()
               else dataclasses.asdict(cfg))
    summary['debug_run'] = False
    for split in ('val', 'test'):
        path = Path(cfg.pred_dir) / f"T00_seed{cfg.seed}_{split}.csv"
        pred = ev.read_pred(str(path))
        metric = ev.compute_metrics(pred.y_true, pred.y_pred, pred.probs)
        summary.update({f"{split}_{k}": metric[k] for k in ('macro_f1', 'top1', 'ece', 'balanced_acc')})
    summary_path.write_text(json.dumps(summary, indent=2), encoding='utf-8')


def run_final(selection_path, baseline, seeds=(0, 1, 2)):
    """Test once per seed for baseline and final; never overwrite test predictions."""
    if len(seeds) < 3 or len(set(seeds)) != len(seeds):
        raise ValueError("Final evaluation requires at least three distinct seeds")
    ev = _import_eval()
    locked = json.loads(Path(selection_path).read_text(encoding="utf-8"))
    recipe, method = Config(**locked["config"]), locked["method"]
    if recipe.debug_samples_per_class is not None or baseline.debug_samples_per_class is not None:
        raise ValueError("Final evaluation cannot use debug data")
    for seed in seeds:
        b = dataclasses.replace(baseline, exp_id="T00", seed=seed, save_test_predictions=True)
        bp = Path(b.pred_dir) / f"T00_seed{seed}_test.csv"
        if not bp.exists():
            # A completed selection run without test output can supply its checkpoint.
            summary = run_dir(b) / "summary.json"
            if summary.exists():
                validate_saved_config(b)
                device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
                model, _, _ = load_state(b, device)
                _, _, test = ds.load_split(b.labels_dir, b.fold)
                names, y, probs = predict_method(model, _loader(b, test), device)
                ev.save_predictions(bp, names, y, probs)
                old = json.loads(summary.read_text(encoding="utf-8"))
                m = ev.compute_metrics(y, probs.argmax(1), probs)
                old.update({f"test_{k}": m[k] for k in ("macro_f1", "top1", "ece", "balanced_acc")})
                summary.write_text(json.dumps(old, indent=2), encoding="utf-8")
                del model
            else:
                run(b)
        else:
            validate_saved_config(b)
            print(f"Keep existing test predictions: {bp.name}")
        recover_baseline(b)

        f = dataclasses.replace(recipe, exp_id="F01", seed=seed, save_test_predictions=False)
        fp = Path(f.pred_dir) / f"F01_seed{seed}_test.csv"
        if fp.exists() or (run_dir(f) / "final_test_arrays.npz").exists():
            validate_saved_config(f)
            recover_final(f)
            print(f"Keep existing test predictions: {fp.name}")
            continue
        run_experiments([f])
        device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
        model, _, _ = load_state(f, device)
        _, val, test = ds.load_split(f.labels_dir, f.fold)
        vl = _loader(f, val)
        temperature = 1.
        if method == "temperature":
            _, yv, zv = inf.predict_logits(model, vl, device)
            temperature = inf.fit_temperature(zv, yv)
        nv, yv, pv = predict_method(model, vl, device, method, temperature)
        ev.save_predictions(Path(f.pred_dir) / f"F01_seed{seed}_val.csv", nv, yv, pv)
        (run_dir(f) / "final_inference.json").write_text(
            json.dumps({"method": method, "temperature": temperature}), encoding="utf-8")
        result = predict_method(model, _loader(f, test), device, method, temperature,
                                return_logits=method == "temperature")
        names, y, probs = result[:3]
        arrays = {"names": np.asarray(names, dtype=object), "y": y, "probs": probs}
        if method == "temperature":
            arrays['logits'] = result[3]
        np.savez_compressed(run_dir(f) / "final_test_arrays.npz", **arrays)
        recover_final(f)
        del model
        if torch.cuda.is_available():
            torch.cuda.empty_cache()


def export_results(root):
    """Seven rubric sheets with provenance, genuine logs, and ddof=1 seed summaries."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    root = Path(root)
    rows = []
    for p in sorted((root / "runs").glob("*/seed*/summary.json")):
        r = json.loads(p.read_text(encoding="utf-8"))
        if r.get("debug_run"):
            continue
        r = {k: v for k, v in r.items() if not isinstance(v, (list, dict))}
        r["source"] = str(p.relative_to(root))
        cp = p.parent / "config.json"
        if cp.exists():
            config = json.loads(cp.read_text(encoding="utf-8"))
            for key in ("label_smoothing", "focal_gamma", "mix_alpha", "sampler", "ema_decay", "drop_rate"):
                r.setdefault(key, config.get(key))
        rows.append(r)
    infer, latency = [], []
    for p in sorted((root / "runs/inference").glob("*/seed*/I*.json")):
        r = json.loads(p.read_text(encoding="utf-8"))
        timing = r.pop("latency")
        r.update({f"{k}_ms": timing[k] for k in ("p50", "p95", "p99")})
        r["images_per_s"] = timing["images_per_s"]
        r["source"] = str(p.relative_to(root))
        infer.append(r)
        latency.append({"exp_id": r["exp_id"], "source_exp_id": r["source_exp_id"],
                        "seed": r["seed"], "method": r["method"], "fused_bn": False, **timing})
    backbones = [r for r in rows if r["exp_id"].startswith("B")]
    training = [r.copy() for r in rows if r["exp_id"].startswith("T")]
    baseline_scores = {r["seed"]: r.get("val_macro_f1") for r in training if r["exp_id"] == "T00"}
    for r in training:
        if r["seed"] in baseline_scores:
            r["delta_vs_T00"] = r["val_macro_f1"] - baseline_scores[r["seed"]]
        baseline_config = next((b for b in training if b["exp_id"] == "T00" and b["seed"] == r["seed"]), None)
        if baseline_config:
            keys = ("init", "aug", "loss", "label_smoothing", "focal_gamma", "mix", "sampler",
                    "lr_backbone", "lr_head", "ema_decay")
            r["changed_fields"] = "; ".join(f"{k}={r.get(k)}" for k in keys if r.get(k) != baseline_config.get(k)) or "baseline"
    finals = [r.copy() for r in rows if r["exp_id"] in ("F01", "T00") and "test_macro_f1" in r]
    for exp in ("F01", "T00"):
        group = [r for r in finals if r["exp_id"] == exp]
        if group:
            agg = {"exp_id": exp, "seed": "mean_std", "n_seeds": len(group)}
            for key in ("test_macro_f1", "test_top1", "test_ece", "val_macro_f1"):
                values = [r[key] for r in group if r.get(key) is not None]
                if values:
                    agg[key + "_mean"] = float(np.mean(values))
                    agg[key + "_std"] = float(np.std(values, ddof=1)) if len(values) > 1 else None
            finals.append(agg)
    perclass = []
    ev = _import_eval()
    for exp in ("F01", "T00"):
        for p in sorted((root / "predictions").glob(f"{exp}_seed*_test.csv")):
            pred = ev.read_pred(str(p))
            m = ev.compute_metrics(pred.y_true, pred.y_pred, pred.probs)
            for k, name in enumerate(ds.CLASS_NAMES):
                perclass.append({"exp_id": exp, "seed": pred.seed, "class": name,
                                 "support": int(m["support"][k]),
                                 **{key: float(m[key][k]) for key in ("precision", "recall", "f1")},
                                 "source": str(p.relative_to(root))})
    sheets = {"Summary": sorted(rows + infer, key=lambda r: r.get("val_macro_f1", -1), reverse=True)[:10],
              "Backbones": backbones, "Training": training, "Inference": infer,
              "Final": finals, "PerClass": perclass, "Latency": latency}
    wb = Workbook()
    wb.remove(wb.active)
    for name, data in sheets.items():
        ws = wb.create_sheet(name)
        if not data:
            ws.append(["status"])
            ws.append(["Chưa có kết quả chạy thật cho giai đoạn này"])
        else:
            cols = list(dict.fromkeys(k for r in data for k in r))
            ws.append(cols)
            for r in data:
                ws.append([r.get(k) for k in cols])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 36
        for column in ws.columns:
            letter = get_column_letter(column[0].column)
            ws.column_dimensions[letter].width = min(42, max(15, max(len(str(c.value or "")) for c in column) + 2))
            for cell in column[1:]:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
                cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str))
    path = root / "results.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def smoke(root, profile="local", *, images_dir=None, labels_dir=None):
    base = base_config(root, profile)
    if images_dir is not None:
        base.images_dir = str(images_dir)
    if labels_dir is not None:
        base.labels_dir = str(labels_dir)
    cfg = dataclasses.replace(base, exp_id="SMOKE", backbone="mobilenetv3", init="scratch",
                              epochs=1, batch_size=2, img_size=64, cache_images=True,
                              num_workers=0, debug_samples_per_class=2,
                              out_dir=str(Path(root) / "runs/_smoke"),
                              pred_dir=str(Path(root) / "runs/_smoke/predictions"),
                              logit_dir=str(Path(root) / "runs/_smoke/logits"),
                              curves_dir=str(Path(root) / "curves/_smoke"))
    return run(cfg)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("stage", choices=("smoke", "export"))
    ap.add_argument("--root", default=str(Path(__file__).resolve().parent.parent))
    ap.add_argument("--profile", choices=("local", "kaggle"), default="local")
    args = ap.parse_args()
    result = smoke(args.root, args.profile) if args.stage == "smoke" else export_results(args.root)
    print(json.dumps(result, ensure_ascii=False, indent=2) if isinstance(result, dict) else result)


if __name__ == "__main__":
    main()
